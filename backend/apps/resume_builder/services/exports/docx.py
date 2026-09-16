from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from .security import split_text_and_links

ACCENT = RGBColor(0x24, 0x57, 0xD6)
DARK = RGBColor(0x18, 0x20, 0x33)
MUTED = RGBColor(0x66, 0x70, 0x85)
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def _set_font(style, *, name="Arial", size=10, bold=False, color=DARK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    style.element.rPr.rFonts.set(qn("w:ascii"), name)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_bottom_rule(style):
    p_pr = style.element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "2457D6")
    p_bdr.append(bottom)


def _configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    _set_font(normal, size=9.55)
    normal.paragraph_format.space_after = Pt(3.2)
    normal.paragraph_format.line_spacing = 1.04

    definitions = (
        ("Resume Name", 19, True, ACCENT, 0, 3),
        ("Resume Title", 10.6, False, DARK, 0, 3),
        ("Resume Contact", 8.8, False, MUTED, 0, 1),
        ("Resume Heading", 10.1, True, ACCENT, 6, 4.5),
        ("Resume Body", 9.55, False, DARK, 0, 3.2),
        ("Resume Skill", 9.25, False, DARK, 0, 2),
        ("Resume Entry Heading", 9.65, True, DARK, 2.8, 1.8),
        ("Resume Detail", 8.9, False, MUTED, 0, 1.8),
        ("Resume Bullet", 9.25, False, DARK, 0, 1.9),
    )
    for name, size, bold, color, before, after in definitions:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_font(style, size=size, bold=bold, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.04

    styles["Resume Heading"].paragraph_format.keep_with_next = True
    styles["Resume Entry Heading"].paragraph_format.keep_with_next = True
    _add_bottom_rule(styles["Resume Heading"])
    styles["Resume Detail"].paragraph_format.left_indent = Inches(0.18)
    styles["Resume Bullet"].base_style = styles["List Bullet"]
    styles["Resume Bullet"].paragraph_format.left_indent = Inches(0.2)
    styles["Resume Bullet"].paragraph_format.first_line_indent = Inches(-0.13)


def _add_text(paragraph, text):
    chunks = text.split("\n")
    for index, chunk in enumerate(chunks):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(chunk)


def _add_hyperlink(paragraph, visible, url):
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2457D6")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = visible
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _fill_paragraph(paragraph, value):
    for text, url in split_text_and_links(value):
        if url:
            _add_hyperlink(paragraph, text, url)
        else:
            _add_text(paragraph, text)


def _prune_relationships(data):
    root = etree.fromstring(data)
    for relationship in list(root):
        target = (relationship.get("Target") or "").lower()
        relation_type = (relationship.get("Type") or "").lower()
        if "customxml" in target or "thumbnail" in target or relation_type.endswith("/customxml") or relation_type.endswith("/metadata/thumbnail"):
            root.remove(relationship)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _prune_content_types(data):
    root = etree.fromstring(data)
    for entry in list(root):
        part_name = (entry.get("PartName") or "").lower()
        extension = (entry.get("Extension") or "").lower()
        if part_name.startswith("/customxml/") or "thumbnail" in part_name or extension in {"jpeg", "jpg", "png", "gif"}:
            root.remove(entry)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _deterministic_package(raw):
    source = BytesIO(raw)
    output = BytesIO()
    with ZipFile(source, "r") as archive:
        entries = {}
        for name in archive.namelist():
            lowered = name.lower()
            if lowered.startswith("customxml/") or lowered == "docprops/thumbnail.jpeg":
                continue
            data = archive.read(name)
            if name in {"_rels/.rels", "word/_rels/document.xml.rels"}:
                data = _prune_relationships(data)
            elif name == "[Content_Types].xml":
                data = _prune_content_types(data)
            entries[name] = data
    with ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, entries[name])
    return output.getvalue()


def render_docx(document_model):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.5)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)
    _configure_styles(document)

    properties = document.core_properties
    properties.title = f"{document_model.name} - Software Engineer CV" if document_model.name else "Software Engineer CV"
    properties.subject = f"resume-content-sha256:{document_model.content_hash}"
    properties.author = document_model.name or "Portfolio resume export service"
    properties.last_modified_by = document_model.name or "Portfolio resume export service"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    properties.revision = 1
    properties.keywords = ""
    properties.category = ""
    properties.comments = ""

    identity = (
        (document_model.name.upper() if document_model.name else "", "Resume Name"),
        (document_model.professional_title, "Resume Title"),
        *((contact, "Resume Contact") for contact in document_model.contacts),
    )
    for value, style_name in identity:
        if value:
            paragraph = document.add_paragraph(style=style_name)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _fill_paragraph(paragraph, value)

    style_by_kind = {
        "body": "Resume Body",
        "skill": "Resume Skill",
        "entry_heading": "Resume Entry Heading",
        "detail": "Resume Detail",
        "bullet": "Resume Bullet",
    }

    for section_model in document_model.sections:
        heading = document.add_paragraph(style="Resume Heading")
        heading.add_run(section_model.heading)
        for item in section_model.items:
            paragraph = document.add_paragraph(style=style_by_kind.get(item.kind, "Resume Body"))
            _fill_paragraph(paragraph, item.text)

    buffer = BytesIO()
    document.save(buffer)
    return _deterministic_package(buffer.getvalue())
