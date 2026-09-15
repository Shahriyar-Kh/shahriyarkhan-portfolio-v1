import re
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfReader

from apps.resume_builder.services.exceptions import ExportIntegrityError

from .security import MAX_ARTIFACT_BYTES, validate_url

_RELATIONSHIP_SUFFIXES = {
    "/officeDocument",
    "/core-properties",
    "/extended-properties",
    "/styles",
    "/stylesWithEffects",
    "/settings",
    "/webSettings",
    "/fontTable",
    "/theme",
    "/numbering",
    "/hyperlink",
}


def _parse_xml(data):
    upper = data.upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise ExportIntegrityError("DOCX contains prohibited XML declarations.")
    parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True)
    return etree.fromstring(data, parser=parser)


def _validate_size(artifact):
    if not isinstance(artifact, bytes) or not artifact or len(artifact) > MAX_ARTIFACT_BYTES:
        raise ExportIntegrityError("Export artifact size is invalid.")


def _normalized_text(value):
    return re.sub(r"\s+", " ", value or "").strip()


def _assert_semantic_order(text, document_model):
    if document_model is None:
        return
    haystack = _normalized_text(text)
    cursor = 0
    for line in document_model.semantic_lines:
        marker = _normalized_text(line)
        position = haystack.find(marker, cursor)
        if position < 0:
            raise ExportIntegrityError("Export content or section order is invalid.")
        cursor = position + len(marker)


def _pdf_has_image(resources, visited=None):
    visited = visited or set()
    if not resources:
        return False
    resources = resources.get_object()
    identity = id(resources)
    if identity in visited:
        return False
    visited.add(identity)
    xobjects = resources.get("/XObject")
    if not xobjects:
        return False
    for reference in xobjects.get_object().values():
        obj = reference.get_object()
        if obj.get("/Subtype") == "/Image":
            return True
        if obj.get("/Subtype") == "/Form" and _pdf_has_image(obj.get("/Resources"), visited):
            return True
    return False


def _validate_pdf(artifact, expected_content_hash, document_model):
    if not artifact.startswith(b"%PDF-"):
        raise ExportIntegrityError("PDF signature is invalid.")
    try:
        reader = PdfReader(BytesIO(artifact), strict=True)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 2:
            raise ExportIntegrityError("PDF encryption or page count is invalid.")
        metadata = reader.metadata or {}
        if metadata.get("/Subject") != f"resume-content-sha256:{expected_content_hash}":
            raise ExportIntegrityError("PDF snapshot binding is invalid.")
        root = reader.trailer["/Root"]
        if any(key in root for key in ("/OpenAction", "/AA", "/AcroForm")):
            raise ExportIntegrityError("PDF contains an active action.")
        names = root.get("/Names")
        if names and any(key in names.get_object() for key in ("/JavaScript", "/EmbeddedFiles")):
            raise ExportIntegrityError("PDF contains prohibited embedded content.")
        text_parts = []
        for page in reader.pages:
            if _pdf_has_image(page.get("/Resources")):
                raise ExportIntegrityError("PDF contains an image.")
            text_parts.append(page.extract_text() or "")
            for annotation_ref in page.get("/Annots", []):
                annotation = annotation_ref.get_object()
                action = annotation.get("/A")
                if annotation.get("/Subtype") != "/Link" or not action or action.get("/S") != "/URI":
                    raise ExportIntegrityError("PDF contains an unsafe annotation.")
                validate_url(str(action.get("/URI")))
        text = "\n".join(text_parts)
    except ExportIntegrityError:
        raise
    except Exception as exc:
        raise ExportIntegrityError("PDF structure is invalid.") from exc
    if not _normalized_text(text):
        raise ExportIntegrityError("PDF has no extractable text.")
    _assert_semantic_order(text, document_model)
    return text


def _validate_docx_relationships(archive):
    for name in archive.namelist():
        if not name.endswith(".rels"):
            continue
        root = _parse_xml(archive.read(name))
        for relationship in root:
            relation_type = relationship.get("Type") or ""
            if not any(relation_type.endswith(suffix) for suffix in _RELATIONSHIP_SUFFIXES):
                raise ExportIntegrityError("DOCX contains an unexpected relationship.")
            target = relationship.get("Target") or ""
            if relationship.get("TargetMode") == "External":
                if not relation_type.endswith("/hyperlink"):
                    raise ExportIntegrityError("DOCX contains an unexpected external relationship.")
                validate_url(target)
            elif target.startswith(("/", chr(92))) or ".." in target.split("/"):
                raise ExportIntegrityError("DOCX contains an unsafe internal relationship.")


def _validate_docx(artifact, expected_content_hash, document_model):
    if not artifact.startswith(b"PK"):
        raise ExportIntegrityError("DOCX signature is invalid.")
    try:
        with ZipFile(BytesIO(artifact), "r") as archive:
            infos = archive.infolist()
            if len(infos) > 100 or len({info.filename for info in infos}) != len(infos):
                raise ExportIntegrityError("DOCX archive entries are invalid.")
            if sum(info.file_size for info in infos) > 25 * 1024 * 1024:
                raise ExportIntegrityError("DOCX expanded size is invalid.")
            if any(info.file_size > 10 * 1024 * 1024 for info in infos):
                raise ExportIntegrityError("DOCX package part is too large.")
            if any(info.filename.startswith(("/", chr(92))) or ".." in info.filename.split("/") for info in infos):
                raise ExportIntegrityError("DOCX archive path is unsafe.")
            for info in infos:
                if info.filename.lower().endswith((".xml", ".rels")):
                    upper = archive.read(info.filename).upper()
                    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                        raise ExportIntegrityError("DOCX contains prohibited XML declarations.")
            if archive.testzip() is not None:
                raise ExportIntegrityError("DOCX archive is corrupt.")
            names = {name.lower() for name in archive.namelist()}
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                raise ExportIntegrityError("DOCX package is incomplete.")
            forbidden = ("word/media/", "vbaproject", "macros", "word/embeddings/", "word/activex/", "customxml/", "thumbnail")
            if any(any(token in name for token in forbidden) for name in names):
                raise ExportIntegrityError("DOCX contains prohibited embedded content.")
            document_xml = _parse_xml(archive.read("word/document.xml"))
            prohibited_tags = {"tbl", "drawing", "pict", "object", "altChunk"}
            if any(etree.QName(element).localname in prohibited_tags for element in document_xml.iter()):
                raise ExportIntegrityError("DOCX contains prohibited layout or embedded content.")
            _validate_docx_relationships(archive)

        parsed = Document(BytesIO(artifact))
        if parsed.tables or parsed.inline_shapes:
            raise ExportIntegrityError("DOCX contains tables or images.")
        if parsed.core_properties.subject != f"resume-content-sha256:{expected_content_hash}":
            raise ExportIntegrityError("DOCX snapshot binding is invalid.")
        text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    except ExportIntegrityError:
        raise
    except Exception as exc:
        raise ExportIntegrityError("DOCX structure is invalid.") from exc
    if not _normalized_text(text):
        raise ExportIntegrityError("DOCX has no extractable text.")
    _assert_semantic_order(text, document_model)
    return text


def validate_artifact(*, format_name, artifact, expected_content_hash, document_model=None):
    _validate_size(artifact)
    if format_name == "pdf":
        return _validate_pdf(artifact, expected_content_hash, document_model)
    if format_name == "docx":
        return _validate_docx(artifact, expected_content_hash, document_model)
    raise ExportIntegrityError("Unsupported export format.")


def extract_artifact_text(*, format_name, artifact):
    _validate_size(artifact)
    try:
        if format_name == "pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(artifact), strict=True).pages)
        if format_name == "docx":
            return "\n".join(paragraph.text for paragraph in Document(BytesIO(artifact)).paragraphs)
    except Exception as exc:
        raise ExportIntegrityError("Export text cannot be extracted.") from exc
    raise ExportIntegrityError("Unsupported export format.")
